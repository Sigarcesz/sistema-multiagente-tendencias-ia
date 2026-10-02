#!/usr/bin/env python3
"""Genera el documento Word final (portada, tabla de contenido, informe, referencias).

GUARDA: se niega a generar el documento si la auditoría no está APROBADO.
Esta regla vive en código, no solo en el prompt: un orquestador confundido
no puede publicar un informe NO_APROBADO.

Uso:
  python3 generar_docx.py --informe runs/<id>/informe_v2.md --auditoria runs/<id>/auditoria_v2.json \
      --hallazgos runs/<id>/03_hallazgos_validados.json --cliente "Horizonte Digital S.A.S." \
      --titulo "Tendencias emergentes de inteligencia artificial..." --salida-dir runs/<id>/final
"""
import argparse
import json
import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

AZUL = RGBColor(0x1F, 0x3A, 0x5F)


def runs_con_negrita(par, texto):
    for i, trozo in enumerate(re.split(r"\*\*(.+?)\*\*", texto)):
        if trozo:
            r = par.add_run(trozo)
            r.bold = i % 2 == 1


def campo_toc(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    for tipo, texto in (("begin", None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ("separate", None), (None, None), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
            r._r.append(el)
        elif texto:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = texto
            r._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = "Clic derecho → Actualizar campo para generar la tabla de contenido."
            r._r.append(t)


def tabla_md(doc, filas):
    celdas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas
              if not re.fullmatch(r"\s*\|?[\s:\-|]+\|?\s*", f)]
    if not celdas:
        return
    ncol = max(len(c) for c in celdas)
    t = doc.add_table(rows=len(celdas), cols=ncol)
    t.style = "Light Grid Accent 1"
    for i, fila in enumerate(celdas):
        for j in range(ncol):
            par = t.cell(i, j).paragraphs[0]
            runs_con_negrita(par, fila[j] if j < len(fila) else "")
            if i == 0:
                for r in par.runs:
                    r.bold = True


def volcar_markdown(doc, md):
    lineas = md.splitlines()
    i = 0
    while i < len(lineas):
        l = lineas[i].rstrip()
        if l.startswith("# "):           # título del informe: ya está en la portada
            pass
        elif l.startswith("## "):
            doc.add_heading(l[3:].strip(), level=1)
        elif l.startswith("### "):
            doc.add_heading(l[4:].strip(), level=2)
        elif l.startswith("|"):
            bloque = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                bloque.append(lineas[i])
                i += 1
            tabla_md(doc, bloque)
            continue
        elif re.match(r"^\s*[-*]\s+", l):
            runs_con_negrita(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", l))
        elif re.match(r"^\s*\d+\.\s+", l):
            runs_con_negrita(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\.\s+", "", l))
        elif l.strip() and not l.strip().startswith("---"):
            runs_con_negrita(doc.add_paragraph(), l.strip())
        i += 1


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--informe", required=True)
    p.add_argument("--auditoria", required=True)
    p.add_argument("--hallazgos", required=True)
    p.add_argument("--cliente", required=True)
    p.add_argument("--titulo", required=True)
    p.add_argument("--salida-dir", required=True)
    a = p.parse_args()

    aud = json.load(open(a.auditoria, encoding="utf-8"))
    if aud.get("estado") != "APROBADO":
        print(f"BLOQUEADO: la auditoría está en estado {aud.get('estado')}. No se genera el documento.", file=sys.stderr)
        sys.exit(3)

    md = open(a.informe, encoding="utf-8").read()
    h = json.load(open(a.hallazgos, encoding="utf-8"))
    hoy = date.today().isoformat()

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    # Portada
    for _ in range(6):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(a.titulo)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(24), AZUL
    for texto, tam in ((f"Informe ejecutivo para {a.cliente}", 14), (f"Fecha: {hoy}", 12),
                       (f"Estado de auditoría: {aud['estado']} · Score global: {aud['score_global']}/100", 11),
                       ("Elaborado por un sistema multiagente: Trend Scout → Report Writer → Quality Auditor & Publisher", 10)):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = par.add_run(texto)
        rr.font.size = Pt(tam)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Tabla de contenido
    doc.add_heading("Tabla de contenido", level=1)
    campo_toc(doc)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Cuerpo
    volcar_markdown(doc, md)

    # Anexo de calidad
    doc.add_heading("Anexo: resultado de la auditoría", level=1)
    filas = ["| Criterio | Puntaje | Umbral |", "|---|---|---|"]
    for c, v in aud["puntajes"].items():
        filas.append(f"| {c} | {v} | {aud['umbrales'].get(c, '—')} |")
    filas.append(f"| **global** | **{aud['score_global']}** | {aud['umbrales']['global']} |")
    tabla_md(doc, filas)

    salida_dir = Path(a.salida_dir)
    salida_dir.mkdir(parents=True, exist_ok=True)
    cliente_slug = re.sub(r"[^A-Za-z0-9]+", "_", a.cliente.split(" S.A.S")[0]).strip("_")
    nombre = f"Informe_Tendencias_IA_{cliente_slug}_{hoy}.docx"
    ruta = salida_dir / nombre
    doc.save(ruta)
    print(json.dumps({"documento": nombre, "ruta": str(ruta), "evidencias_referenciadas": sum(len(t["evidencias"]) for t in h["tendencias"])}, ensure_ascii=False))


if __name__ == "__main__":
    main()
